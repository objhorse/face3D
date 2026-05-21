from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "project_book_assets"
OUT = DOCS / "3D人脸重建与皮肤分析系统_软件项目书_预期工作量与效果.docx"


def font_path():
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for p in candidates:
        if p.exists():
            return str(p)
    return None


FONT = font_path()


def load_font(size, bold=False):
    if FONT:
        try:
            return ImageFont.truetype(FONT, size=size)
        except Exception:
            pass
    return ImageFont.load_default()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(10)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "1F4E79")
        set_cell_text(cell, h, bold=True, color="FFFFFF")
        if widths:
            cell.width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        if level == 1:
            r.font.color.rgb = RGBColor(31, 78, 121)


def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix):])
        runs = [r1, r2]
    else:
        runs = [p.add_run(text)]
    for r in runs:
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        r.font.size = Pt(10.5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.name = "微软雅黑"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        r.font.size = Pt(10.5)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89, 89, 89)


def wrap_text(draw, text, font, max_width):
    lines = []
    current = ""
    for ch in text:
        trial = current + ch
        bbox = draw.textbbox((0, 0), trial, font=font)
        if bbox[2] - bbox[0] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw, rect, text, font, fill=(32, 32, 32), max_lines=3):
    x1, y1, x2, y2 = rect
    lines = wrap_text(draw, text, font, x2 - x1 - 28)[:max_lines]
    line_h = font.size + 8
    total_h = len(lines) * line_h
    y = y1 + (y2 - y1 - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, fill=fill, font=font)
        y += line_h


def arrow(draw, start, end, color=(31, 78, 121), width=4):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        pts = [(x2, y2), (x2 - 16, y2 - 9), (x2 - 16, y2 + 9)]
    else:
        pts = [(x2, y2), (x2 + 16, y2 - 9), (x2 + 16, y2 + 9)]
    draw.polygon(pts, fill=color)


def create_architecture_diagram(path):
    img = Image.new("RGB", (1800, 980), "#F7FAFC")
    draw = ImageDraw.Draw(img)
    title_f = load_font(44)
    h_f = load_font(28)
    t_f = load_font(24)
    draw.text((70, 45), "系统总体架构与预期建设链路", fill="#1F4E79", font=title_f)
    boxes = [
        ((80, 170, 370, 310), "三视图采集\n正脸/左侧/右侧", "#DDEBF7"),
        ((450, 170, 740, 310), "图像预处理\n关键点/人脸分割/校准", "#E2F0D9"),
        ((820, 170, 1110, 310), "3D几何重建\nFLAME/MICA/DECA", "#FFF2CC"),
        ((1190, 170, 1480, 310), "纹理融合\nUV贴图/接缝修复", "#FCE4D6"),
        ((1510, 170, 1760, 310), "3D展示\nWeb交互/报告输出", "#EADCF8"),
    ]
    for rect, text, color in boxes:
        draw.rounded_rectangle(rect, radius=22, fill=color, outline="#6B7280", width=3)
        draw_centered_text(draw, rect, text, h_f)
    for i in range(len(boxes) - 1):
        arrow(draw, (boxes[i][0][2] + 20, 240), (boxes[i + 1][0][0] - 20, 240))
    layer_boxes = [
        ((180, 430, 510, 560), "数据层\n原始照片、会话记录、模型文件、诊断图层"),
        ((590, 430, 920, 560), "算法层\n初始化、拟合优化、深度估计、纹理烘焙"),
        ((1000, 430, 1330, 560), "服务层\nFastAPI接口、任务队列、WebSocket进度"),
        ((1410, 430, 1710, 560), "应用层\n医生查看、客户沟通、术前术后对比"),
    ]
    for rect, text in layer_boxes:
        draw.rounded_rectangle(rect, radius=18, fill="#FFFFFF", outline="#A5B4C5", width=3)
        draw_centered_text(draw, rect, text, t_f)
    notes = [
        "重点：本项目书描述拟建设工作范围和预期效果，图片为功能示意/阶段参考。",
        "建设目标：形成可交付的软件系统，而非单一算法 Demo。",
        "验收关注：流程可跑通、结果可查看、指标可复核、异常可追踪。"
    ]
    y = 690
    for note in notes:
        draw.rounded_rectangle((170, y, 1630, y + 72), radius=14, fill="#FFFFFF", outline="#D0D7DE", width=2)
        draw.text((210, y + 20), note, fill="#374151", font=t_f)
        y += 95
    img.save(path, quality=95)


def create_plan_diagram(path):
    img = Image.new("RGB", (1800, 900), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_f = load_font(42)
    h_f = load_font(26)
    t_f = load_font(22)
    draw.text((70, 45), "建议实施阶段与工作量分布", fill="#1F4E79", font=title_f)
    stages = [
        ("阶段1\n需求确认与方案细化", "接口、硬件、样本、验收口径"),
        ("阶段2\n采集与数据管理", "三视图采集、会话管理、数据归档"),
        ("阶段3\n重建算法开发", "几何初始化、拟合优化、纹理融合"),
        ("阶段4\n皮肤分析与图层", "红区、色沉、纹理/皱纹、报告指标"),
        ("阶段5\n前后端集成", "API、Web展示、进度、模型下载"),
        ("阶段6\n测试部署验收", "样本测试、性能优化、交付文档"),
    ]
    x = 90
    y = 180
    w = 260
    h = 180
    colors = ["#DDEBF7", "#E2F0D9", "#FFF2CC", "#FCE4D6", "#EADCF8", "#D9EAD3"]
    for idx, (name, detail) in enumerate(stages):
        rect = (x + idx * 280, y, x + idx * 280 + w, y + h)
        draw.rounded_rectangle(rect, radius=20, fill=colors[idx], outline="#6B7280", width=3)
        draw_centered_text(draw, (rect[0] + 10, rect[1] + 14, rect[2] - 10, rect[1] + 95), name, h_f, max_lines=3)
        draw_centered_text(draw, (rect[0] + 12, rect[1] + 96, rect[2] - 12, rect[3] - 12), detail, t_f, max_lines=3)
        if idx < len(stages) - 1:
            arrow(draw, (rect[2] + 10, y + 90), (rect[2] + 35, y + 90), width=3)
    bars = [
        ("需求/产品设计", 0.16, "#5B9BD5"),
        ("算法与工程开发", 0.36, "#70AD47"),
        ("前后端与系统集成", 0.22, "#FFC000"),
        ("测试、部署与培训", 0.18, "#ED7D31"),
        ("项目管理与文档", 0.08, "#A5A5A5"),
    ]
    bx, by = 210, 520
    draw.text((90, 470), "预期工作量占比（建议口径）", fill="#1F4E79", font=h_f)
    for label, pct, color in bars:
        draw.text((bx, by + 8), label, fill="#374151", font=t_f)
        draw.rounded_rectangle((560, by, 1500, by + 42), radius=18, fill="#EEF2F7")
        draw.rounded_rectangle((560, by, 560 + int(940 * pct), by + 42), radius=18, fill=color)
        draw.text((1530, by + 8), f"{int(pct * 100)}%", fill="#374151", font=t_f)
        by += 62
    img.save(path, quality=95)


def create_three_view_diagram(path):
    img = Image.new("RGB", (1600, 760), "#F7FAFC")
    draw = ImageDraw.Draw(img)
    title_f = load_font(40)
    h_f = load_font(28)
    t_f = load_font(22)
    draw.text((60, 40), "三视图采集与重建输入示意", fill="#1F4E79", font=title_f)
    centers = [(360, 360), (800, 360), (1240, 360)]
    labels = ["左侧约45度", "正脸", "右侧约45度"]
    for i, (cx, cy) in enumerate(centers):
        draw.ellipse((cx - 90, cy - 120, cx + 90, cy + 120), fill="#F4D8C7", outline="#8E6B5B", width=4)
        draw.ellipse((cx - 45, cy - 35, cx - 18, cy - 8), fill="#FFFFFF", outline="#4B5563", width=2)
        draw.ellipse((cx + 18, cy - 35, cx + 45, cy - 8), fill="#FFFFFF", outline="#4B5563", width=2)
        draw.polygon([(cx, cy - 10), (cx - 18, cy + 40), (cx + 18, cy + 40)], fill="#C99C89")
        draw.arc((cx - 45, cy + 45, cx + 45, cy + 95), 10, 170, fill="#8E4B4B", width=4)
        if i == 0:
            draw.rectangle((cx - 98, cy - 120, cx - 25, cy + 120), fill="#F7FAFC")
            draw.line((cx - 20, cy - 115, cx - 20, cy + 115), fill="#8E6B5B", width=4)
        if i == 2:
            draw.rectangle((cx + 25, cy - 120, cx + 98, cy + 120), fill="#F7FAFC")
            draw.line((cx + 20, cy - 115, cx + 20, cy + 115), fill="#8E6B5B", width=4)
        draw.text((cx - 78, cy + 155), labels[i], fill="#374151", font=h_f)
    for x in [580, 1020]:
        arrow(draw, (x, 360), (x + 110, 360), width=5)
    draw.rounded_rectangle((180, 625, 1420, 700), radius=18, fill="#FFFFFF", outline="#D0D7DE", width=2)
    draw_centered_text(draw, (190, 625, 1410, 700), "预期建设要求：固定光照、固定距离、统一背景、自动质检；输入质量直接影响重建精度。", t_f)
    img.save(path, quality=95)


def safe_copy_resize(src, dst, max_size=(1400, 1000)):
    src = Path(src)
    if not src.exists():
        return None
    img = Image.open(src).convert("RGB")
    img.thumbnail(max_size, Image.LANCZOS)
    img.save(dst, quality=92)
    return dst


def create_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)
    create_architecture_diagram(ASSETS / "architecture.png")
    create_plan_diagram(ASSETS / "plan.png")
    create_three_view_diagram(ASSETS / "three_view.png")
    image_map = {
        "render_expected.png": ROOT / "image.png",
        "geometry_expected.png": ROOT / "image12.png",
        "landmark_reproj.png": ROOT / "output" / "debug" / "pipeline_audit" / "02_geometry" / "landmark_reproj_front.png",
        "uv_texture.png": ROOT / "output" / "sessions" / "34" / "textures" / "albedo_white.png",
    }
    resized = {}
    for name, src in image_map.items():
        dst = ASSETS / name
        out = safe_copy_resize(src, dst)
        if out:
            resized[name] = out
    return resized


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("3D人脸重建与皮肤分析系统")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("软件项目书")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(22)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("拟建设工作范围、预期工作量与预期效果说明")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"编制日期：{date.today().isoformat()}")
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(11)

    doc.add_paragraph()
    note = (
        "重要说明：本文档用于向甲方说明项目拟建设内容、预计投入工作量、阶段性交付物和预期可达到的效果。"
        "文中图片用于功能示意和目标效果参考，不等同于全部功能已经完成，也不作为最终验收唯一依据。"
        "最终实施范围、周期、验收标准和商务条件应以双方确认的合同、需求规格说明书和验收清单为准。"
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(1.2)
    p.paragraph_format.right_indent = Cm(1.2)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(note)
    r.font.name = "微软雅黑"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(89, 89, 89)
    doc.add_page_break()


def build_doc():
    images = create_assets()
    doc = Document()
    configure_document(doc)
    cover(doc)

    add_heading(doc, "一、项目概述", 1)
    add_para(doc, "本项目拟建设一套面向医美、皮肤管理、数字化面诊和术前术后对比场景的 3D 人脸重建与皮肤分析软件系统。系统以正脸、左侧、右侧三视图照片为主要输入，经过图像预处理、几何重建、纹理融合、皮肤图层分析和 Web 端可视化展示，形成可查看、可比对、可归档的 3D 人脸模型和分析结果。")
    add_para(doc, "本项目书重点说明计划建设的工作内容、建议技术路线、预期交付成果、阶段性工作量和验收方式，便于甲方评估项目实施范围和投入强度。")
    add_bullets(doc, [
        "建设定位：从算法验证升级为可交付的软件系统，覆盖采集、计算、展示、管理和报告输出链路。",
        "目标用户：医美机构、皮肤管理机构、咨询师、医生及内部运营人员。",
        "业务价值：提升面诊沟通效率，沉淀客户面部数据，为术前评估、术后对比和皮肤问题跟踪提供数字化依据。",
        "交付原则：先保证稳定可用和可验收，再逐步提升高保真重建、皮肤分析精度和自动化程度。"
    ])
    if (ASSETS / "architecture.png").exists():
        doc.add_picture(str(ASSETS / "architecture.png"), width=Cm(16.5))
        add_caption(doc, "图1 系统总体架构与预期建设链路示意")

    add_heading(doc, "二、建设目标与预期效果", 1)
    add_para(doc, "系统建设完成后，预期实现从三视图照片到 3D 模型、纹理贴图、皮肤分析图层和前端交互展示的完整流程。由于 3D 重建效果受采集设备、光照、拍摄姿态、样本质量和算法模型权重影响，建议将预期效果分为基础可用、优化增强和高保真扩展三个层级进行验收。")
    add_table(doc, ["层级", "预期效果", "说明"], [
        ["基础可用", "完成三视图上传、后台重建、GLB 模型生成、Web 端查看和任务进度反馈", "作为第一阶段验收底线，关注流程完整性和稳定性。"],
        ["优化增强", "提升脸型、鼻部、下颌线、侧脸轮廓和纹理接缝质量，支持调试图和质量评估", "适合正式面诊演示，关注模型真实性和可解释性。"],
        ["高保真扩展", "引入 MICA/DECA/EMOCA 等更强初始化模型，完善红区、色沉、皱纹等皮肤分析图层", "适合作为后续增强建设内容，关注算法效果和临床/业务可用性。"],
    ], widths=[3.0, 7.0, 6.0])
    if images.get("render_expected.png"):
        doc.add_picture(str(images["render_expected.png"]), width=Cm(12.5))
        add_caption(doc, "图2 3D 人脸模型交互展示预期效果参考")
    if images.get("geometry_expected.png"):
        doc.add_picture(str(images["geometry_expected.png"]), width=Cm(12.5))
        add_caption(doc, "图3 几何灰模细节展示预期效果参考")

    add_heading(doc, "三、拟建设功能范围", 1)
    add_table(doc, ["功能模块", "拟建设内容", "预期交付物"], [
        ["三视图采集与导入", "支持正脸、左侧、右侧照片导入；建议增加采集质量提示、角度提示、光照提示和缺失检查。", "图片上传页面、采集规范、输入质量校验逻辑。"],
        ["数据预处理", "完成图像缩放、裁切、关键点检测、人脸区域分割、背景处理和调试图输出。", "预处理脚本、关键点图、Mask 图、错误提示。"],
        ["3D 几何重建", "基于 FLAME/3DMM 完成共享身份形状拟合；接入 MICA/DECA 等初始化能力，提升侧脸、鼻梁、下颌和轮廓真实性。", "OBJ/GLB 模型、相机参数、重投影调试图、初始化参数文件。"],
        ["纹理融合与修复", "将多视角照片映射到 UV 空间，处理亮度色差、接缝、孔洞和不可见区域，生成 2K 级纹理贴图。", "UV 纹理图、来源视图图、置信度图、GLB 材质。"],
        ["皮肤分析图层", "拟实现红区、色沉、纹理/皱纹等分析图层，并支持与 3D 模型或二维报告联动。", "分析图层、指标说明、区域标注和报告字段。"],
        ["Web 端展示", "提供模型旋转、缩放、平移、图层切换、任务列表、模型下载和客户档案查看。", "前端页面、Three.js 查看器、交互控件。"],
        ["后台服务与接口", "提供会话创建、进度推送、模型下载、任务状态查询、数据归档和异常处理接口。", "FastAPI 服务、数据库表、接口文档。"],
        ["测试与部署", "完成样本测试、性能测试、异常流程测试、部署脚本和使用说明。", "测试报告、部署文档、运维说明、验收材料。"],
    ], widths=[3.0, 8.0, 5.0])

    if (ASSETS / "three_view.png").exists():
        doc.add_picture(str(ASSETS / "three_view.png"), width=Cm(15.5))
        add_caption(doc, "图4 三视图采集与重建输入示意")
    if images.get("landmark_reproj.png"):
        doc.add_picture(str(images["landmark_reproj.png"]), width=Cm(12.5))
        add_caption(doc, "图5 关键点重投影与几何拟合质量检查示意")
    if images.get("uv_texture.png"):
        doc.add_picture(str(images["uv_texture.png"]), width=Cm(11.5))
        add_caption(doc, "图6 UV 纹理贴图与多视角融合效果参考")

    add_heading(doc, "四、技术路线建议", 1)
    add_para(doc, "为兼顾可交付性与后续效果提升，建议采用“工程主链路先稳定、算法能力分阶段增强”的技术路线。第一阶段建立可运行的端到端系统，第二阶段针对几何初始化、纹理融合和皮肤分析逐项优化，第三阶段再根据样本效果决定是否引入更强的 dense reconstruction 或光度优化方案。")
    add_table(doc, ["技术环节", "建议方案", "甲方可见价值"], [
        ["采集规范", "固定相机距离、光照、背景、角度和分辨率，必要时提供采集引导界面。", "减少重建失败率，让结果更稳定。"],
        ["关键点与分割", "使用 MediaPipe Face Mesh、面部轮廓 Mask、质量检测和调试图。", "发现照片不合格原因，降低人工排查成本。"],
        ["身份形状初始化", "优先建设 MICA + DECA/EMOCA 初始化后端，并保留 legacy fallback。", "提升脸型、侧脸、鼻部和下颌线真实性。"],
        ["多视图优化", "共享 identity shape，按视图优化 pose/expression/camera，输出重投影误差。", "让模型效果有量化依据。"],
        ["纹理处理", "按可见性、法线夹角和视图置信度融合纹理，使用 inpaint/Poisson/颜色匹配修复接缝。", "减少明显拼接、黑洞、色差和拉伸。"],
        ["前端渲染", "Three.js/GLB 展示，支持纹理、灰模、分析图层切换。", "便于面诊演示和客户沟通。"],
        ["服务接口", "FastAPI + SQLite/可扩展数据库 + WebSocket 进度推送。", "方便与现有业务系统集成。"],
    ], widths=[3.2, 8.0, 4.8])

    add_heading(doc, "五、预期工作量拆分", 1)
    add_para(doc, "以下工作量为项目实施建议口径，实际周期需结合甲方确认的部署环境、硬件设备、样本数量、算法精度要求、是否需要医美业务系统对接等因素进一步评估。")
    add_table(doc, ["阶段", "主要工作", "建议工作量", "关键交付物"], [
        ["阶段1：需求确认与原型设计", "确认采集流程、用户角色、展示页面、报告内容、验收指标和部署环境。", "约 1-2 周", "需求规格说明、原型图、验收清单初稿。"],
        ["阶段2：采集与数据管理", "建设图片上传、会话管理、文件归档、输入校验和质量提示。", "约 1-2 周", "采集端/上传端、数据库结构、数据目录规范。"],
        ["阶段3：重建算法主链路", "完善预处理、几何重建、纹理融合、GLB 导出和调试输出。", "约 3-5 周", "模型生成服务、调试图、GLB 文件、算法参数说明。"],
        ["阶段4：皮肤分析图层", "实现红区、色沉、纹理/皱纹等基础图层，建立指标展示和报告字段。", "约 2-4 周", "分析图层、指标说明、报告页面。"],
        ["阶段5：前后端集成", "建设 Web 查看器、任务进度、模型下载、历史记录和异常提示。", "约 2-3 周", "可演示系统、接口文档、页面联调记录。"],
        ["阶段6：测试部署与验收", "完成样本测试、性能优化、部署、培训和验收材料整理。", "约 1-2 周", "测试报告、部署文档、操作手册、验收报告。"],
    ], widths=[3.0, 7.0, 2.5, 4.0])
    if (ASSETS / "plan.png").exists():
        doc.add_picture(str(ASSETS / "plan.png"), width=Cm(16.0))
        add_caption(doc, "图7 建议实施阶段与预期工作量分布")

    add_heading(doc, "六、交付成果清单", 1)
    add_table(doc, ["类别", "交付内容", "验收方式"], [
        ["软件系统", "前端页面、后端服务、数据库、任务处理流程、模型查看器。", "在测试环境或甲方指定环境完成演示。"],
        ["算法模块", "预处理、几何拟合、纹理融合、皮肤分析、模型导出。", "使用双方确认样本跑通并输出结果。"],
        ["数据与模型文件", "GLB/OBJ、UV 纹理、相机参数、调试图、分析图层。", "检查文件完整性、可打开性和命名规范。"],
        ["接口文档", "上传、查询、下载、进度、会话管理等接口说明。", "按接口清单联调验证。"],
        ["部署文档", "环境依赖、模型权重、启动方式、目录说明、常见问题。", "按文档完成一次独立部署或复核。"],
        ["测试与验收材料", "功能测试、样本测试、异常测试、性能记录和验收报告。", "双方确认问题清单和验收结论。"],
    ], widths=[3.0, 8.5, 4.5])

    add_heading(doc, "七、验收标准建议", 1)
    add_para(doc, "建议将验收分为功能验收、效果验收、性能验收和文档验收四类。对于算法效果，应避免只用单张图片主观判断，建议固定测试样本、固定拍摄规范、固定指标和人工复核流程。")
    add_table(doc, ["验收项", "建议标准", "备注"], [
        ["流程完整性", "三视图上传后可自动完成处理并生成可查看模型。", "允许失败样本给出明确错误原因。"],
        ["模型可视化", "Web 端支持旋转、缩放、平移、模型加载和基础图层切换。", "需兼容主流浏览器。"],
        ["几何质量", "正脸和侧脸轮廓基本一致，关键点重投影误差可输出并可复核。", "高保真指标需按样本集单独确认。"],
        ["纹理质量", "主要面部区域无明显大面积黑洞、错位和严重色差。", "耳部、发际线、遮挡区域可作为低置信区域说明。"],
        ["皮肤图层", "红区、色沉、纹理/皱纹图层可生成、可查看、可解释。", "医学结论需谨慎表述，建议定位为辅助分析。"],
        ["性能", "单人重建耗时达到双方确认目标，例如 1-3 分钟内完成基础结果。", "依赖 GPU、模型规模和图片分辨率。"],
        ["稳定性", "连续测试样本无服务崩溃，异常任务不影响后续任务。", "需记录错误日志。"],
        ["文档", "提供部署、使用、接口和运维说明。", "便于甲方后续维护和二次开发。"],
    ], widths=[3.0, 8.0, 5.0])

    add_heading(doc, "八、甲方配合事项", 1)
    add_bullets(doc, [
        "提供目标应用场景和业务流程，例如面诊、术前术后对比、皮肤管理或内部研究。",
        "确认采集设备、拍摄距离、光照环境、背景颜色、照片分辨率和样本隐私合规要求。",
        "提供一定数量的测试样本，用于算法调试、效果评估和验收。",
        "确认部署环境，包括 GPU/CPU、操作系统、内网/公网访问、数据存储和备份策略。",
        "指定甲方验收人员，参与需求确认、阶段评审、样本效果复核和最终验收。"
    ])

    add_heading(doc, "九、风险与边界说明", 1)
    add_table(doc, ["风险点", "影响", "建议应对"], [
        ["采集质量不稳定", "姿态、光照、遮挡会直接影响重建效果。", "制定采集规范并增加质量检测。"],
        ["算法模型依赖", "MICA/DECA/Depth 等外部模型可能涉及权重、环境和许可问题。", "提前确认模型来源、部署方式和授权边界。"],
        ["高保真预期过高", "照片级纹理不等于医学级测量，部分区域存在不可见和遮挡。", "分级验收，区分展示效果和量化指标。"],
        ["皮肤分析医学边界", "红区、色沉、皱纹分析可作为辅助，不宜直接替代医生诊断。", "文案定位为辅助评估，输出置信度和说明。"],
        ["算力和性能", "GPU 规格不足会导致处理时间变长。", "在部署前完成硬件评估和性能测试。"],
        ["隐私与数据安全", "人脸照片和模型属于敏感数据。", "限定访问权限、存储策略和数据脱敏/删除机制。"],
    ], widths=[3.2, 6.0, 6.8])

    add_heading(doc, "十、建议下一步", 1)
    add_para(doc, "建议在正式立项前安排一次需求确认会，围绕使用场景、验收样本、采集设备、部署环境和效果优先级进行确认。确认后可输出正式《需求规格说明书》《实施计划》《验收标准表》，再进入开发和阶段性交付。")
    add_bullets(doc, [
        "确认第一版是否以“可演示系统”为主，还是以“高保真算法效果”为主。",
        "确认是否需要接入现有客户管理系统、报告系统或门店设备。",
        "确认样本数据数量、授权方式和隐私处理要求。",
        "确认项目周期、阶段验收节点和最终交付形式。"
    ])

    doc.add_page_break()
    add_heading(doc, "附录：文档图片说明", 1)
    add_table(doc, ["图片", "用途", "说明"], [
        ["系统架构图", "说明拟建设软件链路", "由项目方案生成，用于甲方理解整体范围。"],
        ["三视图采集图", "说明输入规范", "为示意图，不代表最终采集设备形态。"],
        ["3D 模型渲染图", "说明预期展示效果", "取自项目本地素材，作为效果参考。"],
        ["灰模细节图", "说明几何重建目标", "用于表达轮廓和细节方向。"],
        ["关键点重投影图", "说明算法质检方式", "用于表达可复核、可调试的验收思路。"],
        ["UV 纹理图", "说明纹理融合工作", "用于表达贴图和接缝修复工作量。"],
    ], widths=[4.0, 5.0, 7.0])

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_doc()
    print(out)
