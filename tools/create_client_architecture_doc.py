from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "client_architecture_assets"
ZYZ_DIR = ROOT / "new_captures" / "zyz_captures"
OUT = DOCS / "3D人脸重建与皮肤分析系统_软件架构书_甲方版_zyz数据集.docx"


def font_path():
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return str(path)
    return None


FONT = font_path()


def load_font(size, bold=False):
    if FONT:
        try:
            return ImageFont.truetype(FONT, size=size)
        except Exception:
            pass
    return ImageFont.load_default()


def set_run_font(run, size=10.5, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, align_center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center or bold else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    set_run_font(run, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "1F4E79")
        set_cell_text(cell, header, bold=True, color="FFFFFF", align_center=True)
        if widths:
            cell.width = Cm(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            if widths:
                cells[idx].width = Cm(widths[idx])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_run_font(run, size=15 if level == 1 else 12.5, bold=True)
        if level == 1:
            run.font.color.rgb = RGBColor(31, 78, 121)
    return paragraph


def add_para(doc, text, first_line=True):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.space_after = Pt(4)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.line_spacing = 1.22
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5)


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color="666666")
    run.italic = True


def wrap_text(draw, text, font, max_width):
    lines = []
    current = ""
    for char in text:
        trial = current + char
        bbox = draw.textbbox((0, 0), trial, font=font)
        if bbox[2] - bbox[0] <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def draw_centered_text(draw, rect, text, font, fill="#222222", max_lines=4):
    x1, y1, x2, y2 = rect
    lines = []
    for part in text.split("\n"):
        lines.extend(wrap_text(draw, part, font, x2 - x1 - 30))
    lines = lines[:max_lines]
    line_height = font.size + 8
    total_height = line_height * len(lines)
    y = y1 + (y2 - y1 - total_height) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, fill=fill, font=font)
        y += line_height


def arrow(draw, start, end, color="#1F4E79", width=5):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        if x2 >= x1:
            points = [(x2, y2), (x2 - 18, y2 - 10), (x2 - 18, y2 + 10)]
        else:
            points = [(x2, y2), (x2 + 18, y2 - 10), (x2 + 18, y2 + 10)]
    else:
        if y2 >= y1:
            points = [(x2, y2), (x2 - 10, y2 - 18), (x2 + 10, y2 - 18)]
        else:
            points = [(x2, y2), (x2 - 10, y2 + 18), (x2 + 10, y2 + 18)]
    draw.polygon(points, fill=color)


def find_zyz_images():
    files = sorted(ZYZ_DIR.glob("camera*.jpg"))
    if len(files) < 3:
        raise FileNotFoundError(f"zyz数据集照片不足3张: {ZYZ_DIR}")
    return files[:3]


def create_zyz_contact_sheet(path):
    images = find_zyz_images()
    canvas = Image.new("RGB", (1800, 880), "#F7FAFC")
    draw = ImageDraw.Draw(canvas)
    title_font = load_font(44)
    label_font = load_font(26)
    small_font = load_font(22)
    draw.text((70, 45), "zyz 数据集三视图采集样例", fill="#1F4E79", font=title_font)
    labels = ["camera1 原始视角", "camera2 原始视角", "camera3 原始视角"]
    x_positions = [90, 640, 1190]
    for image_path, label, x in zip(images, labels, x_positions):
        img = Image.open(image_path).convert("RGB")
        img.thumbnail((470, 590), Image.LANCZOS)
        frame = Image.new("RGB", (500, 620), "#FFFFFF")
        fx = (500 - img.width) // 2
        fy = (620 - img.height) // 2
        frame.paste(img, (fx, fy))
        canvas.paste(frame, (x, 145))
        draw.rounded_rectangle((x, 145, x + 500, 765), radius=18, outline="#CBD5E1", width=4)
        draw.text((x + 28, 785), label, fill="#111827", font=label_font)
        draw.text((x + 28, 823), image_path.name, fill="#64748B", font=small_font)
    canvas.save(path, quality=94)


def create_architecture_diagram(path):
    img = Image.new("RGB", (1800, 980), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42)
    box_font = load_font(25)
    small_font = load_font(21)
    draw.text((70, 45), "系统总体架构：采集、重建、分析、展示一体化", fill="#1F4E79", font=title_font)

    top = [
        ((80, 160, 350, 300), "三视图采集\nzyz 数据集", "#DBEAFE"),
        ((430, 160, 700, 300), "预处理\n人脸检测/对齐/掩膜", "#DCFCE7"),
        ((780, 160, 1050, 300), "高保真几何\nMICA/DECA/FLAME", "#FEF3C7"),
        ((1130, 160, 1400, 300), "纹理融合\nUV 贴图/补洞", "#FFEDD5"),
        ((1480, 160, 1740, 300), "交付输出\nGLB/OBJ/Web预览", "#E0E7FF"),
    ]
    for rect, text, color in top:
        draw.rounded_rectangle(rect, radius=20, fill=color, outline="#64748B", width=3)
        draw_centered_text(draw, rect, text, box_font, max_lines=3)
    for idx in range(len(top) - 1):
        arrow(draw, (top[idx][0][2] + 18, 230), (top[idx + 1][0][0] - 18, 230))

    layers = [
        ((150, 430, 520, 575), "数据层\nzyz三视图照片、相机标定、任务记录、模型文件", "#FFFFFF"),
        ((590, 430, 960, 575), "算法层\n关键点、三维形状、深度增强、纹理融合、质量评估", "#FFFFFF"),
        ((1030, 430, 1400, 575), "服务层\nFastAPI、SQLite、WebSocket进度、任务管理", "#FFFFFF"),
        ((1470, 430, 1740, 575), "展示层\nThree.js模型查看、结果下载、阶段报告", "#FFFFFF"),
    ]
    for rect, text, color in layers:
        draw.rounded_rectangle(rect, radius=18, fill=color, outline="#94A3B8", width=3)
        draw_centered_text(draw, rect, text, small_font, max_lines=4)

    notes = [
        "面向甲方的核心价值：把多视角人脸照片转化为可浏览、可归档、可复核的三维人脸资产。",
        "预期工作内容：完成采集流程、重建算法、皮肤纹理处理、Web展示、结果管理和验收指标闭环。",
        "本架构书中的数据样例统一使用 zyz 数据集，便于甲方看到从真实输入到预期输出的完整链路。",
    ]
    y = 700
    for note in notes:
        draw.rounded_rectangle((150, y, 1650, y + 68), radius=14, fill="#FFFFFF", outline="#CBD5E1", width=2)
        draw.text((190, y + 19), note, fill="#334155", font=small_font)
        y += 88
    img.save(path, quality=94)


def create_data_flow_diagram(path):
    img = Image.new("RGB", (1800, 900), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42)
    box_font = load_font(24)
    small_font = load_font(21)
    draw.text((70, 45), "zyz 数据驱动的处理流程", fill="#1F4E79", font=title_font)

    stages = [
        ("1. 数据导入", "camera1/2/3\n三张原始照片"),
        ("2. 质量检查", "清晰度、曝光、姿态\n与人脸可见范围"),
        ("3. 预处理", "裁剪、标准化、关键点\n人脸区域掩膜"),
        ("4. 几何重建", "MICA身份形状\nDECA姿态表情\nFLAME多视图优化"),
        ("5. 纹理与皮肤", "多视图颜色融合\nUV纹理、皮肤区域特征"),
        ("6. 结果交付", "GLB/OBJ模型\n截图、指标、报告"),
    ]
    x = 95
    y = 180
    w = 245
    h = 190
    colors = ["#DBEAFE", "#E0F2FE", "#DCFCE7", "#FEF3C7", "#FFEDD5", "#E0E7FF"]
    for idx, (title, body) in enumerate(stages):
        rect = (x + idx * 280, y, x + idx * 280 + w, y + h)
        draw.rounded_rectangle(rect, radius=20, fill=colors[idx], outline="#64748B", width=3)
        draw_centered_text(draw, (rect[0] + 8, rect[1] + 18, rect[2] - 8, rect[1] + 78), title, box_font, max_lines=2)
        draw_centered_text(draw, (rect[0] + 12, rect[1] + 86, rect[2] - 12, rect[3] - 12), body, small_font, max_lines=4)
        if idx < len(stages) - 1:
            arrow(draw, (rect[2] + 10, y + 95), (rect[2] + 35, y + 95), width=4)

    draw.text((120, 465), "预期可验收产物", fill="#1F4E79", font=box_font)
    rows = [
        ("输入记录", "zyz 数据集三视图照片、采集时间、视角标注"),
        ("过程数据", "关键点图、掩膜图、深度图、重投影误差、阶段日志"),
        ("三维资产", "带纹理的 GLB/OBJ 模型、UV 贴图、相机参数文件"),
        ("甲方材料", "架构书、阶段进展图、演示页面、验收指标表"),
    ]
    y = 525
    for idx, (name, body) in enumerate(rows):
        draw.rounded_rectangle((160, y, 1640, y + 68), radius=14, fill="#F8FAFC", outline="#CBD5E1", width=2)
        draw.text((205, y + 19), name, fill="#0F172A", font=small_font)
        draw.text((430, y + 19), body, fill="#334155", font=small_font)
        y += 85
    img.save(path, quality=94)


def create_work_plan_diagram(path):
    img = Image.new("RGB", (1800, 820), "#F8FAFC")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42)
    label_font = load_font(24)
    small_font = load_font(20)
    draw.text((70, 45), "预期工作计划与交付节奏", fill="#1F4E79", font=title_font)
    phases = [
        ("阶段一\n需求与采集确认", "明确采集条件、zyz样例数据\n完成输入规范和验收口径"),
        ("阶段二\n核心重建链路", "打通预处理、几何重建\n输出基础三维模型"),
        ("阶段三\n纹理与皮肤分析", "完成UV纹理融合\n沉淀皮肤区域分析结果"),
        ("阶段四\n系统集成展示", "接入API、任务管理\nWeb端预览和结果下载"),
        ("阶段五\n验收与移交", "输出文档、测试报告\n部署说明和维护建议"),
    ]
    x0 = 120
    y0 = 190
    w = 300
    h = 160
    for idx, (title, body) in enumerate(phases):
        x = x0 + idx * 330
        rect = (x, y0, x + w, y0 + h)
        draw.rounded_rectangle(rect, radius=20, fill="#FFFFFF", outline="#64748B", width=3)
        draw_centered_text(draw, (x + 10, y0 + 12, x + w - 10, y0 + 78), title, label_font, max_lines=3)
        draw_centered_text(draw, (x + 14, y0 + 82, x + w - 14, y0 + h - 12), body, small_font, max_lines=4)
        if idx < len(phases) - 1:
            arrow(draw, (x + w + 8, y0 + h // 2), (x + w + 28, y0 + h // 2), width=4)

    bars = [
        ("采集与数据治理", 0.16, "#2563EB"),
        ("重建算法与质量优化", 0.34, "#16A34A"),
        ("皮肤纹理与分析能力", 0.20, "#F59E0B"),
        ("系统集成与展示", 0.18, "#7C3AED"),
        ("测试验收与文档移交", 0.12, "#64748B"),
    ]
    draw.text((120, 460), "预期工作量构成", fill="#1F4E79", font=label_font)
    y = 525
    for label, pct, color in bars:
        draw.text((160, y + 7), label, fill="#334155", font=small_font)
        draw.rounded_rectangle((520, y, 1500, y + 38), radius=16, fill="#E2E8F0")
        draw.rounded_rectangle((520, y, 520 + int(980 * pct), y + 38), radius=16, fill=color)
        draw.text((1530, y + 7), f"{int(pct * 100)}%", fill="#334155", font=small_font)
        y += 58
    img.save(path, quality=94)


def create_assets():
    ASSETS.mkdir(parents=True, exist_ok=True)
    create_zyz_contact_sheet(ASSETS / "zyz_three_view.png")
    create_architecture_diagram(ASSETS / "client_architecture.png")
    create_data_flow_diagram(ASSETS / "zyz_data_flow.png")
    create_work_plan_diagram(ASSETS / "work_plan.png")


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = doc.styles
    styles["Normal"].font.name = "微软雅黑"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "微软雅黑"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def add_cover(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("3D人脸重建与皮肤分析系统")
    set_run_font(run, size=26, bold=True, color="1F4E79")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("软件架构书")
    set_run_font(run, size=22, bold=True)

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = desc.add_run("甲方沟通版：预期工作内容、系统架构与交付说明")
    set_run_font(run, size=13, color="666666")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"编制日期：{date.today().isoformat()}")
    set_run_font(run, size=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("数据样例：zyz 三视图采集数据集")
    set_run_font(run, size=11)

    doc.add_paragraph()
    doc.add_picture(str(ASSETS / "zyz_three_view.png"), width=Cm(16.2))
    add_caption(doc, "图 1  zyz 数据集三视图采集样例")
    doc.add_page_break()


def build_doc():
    create_assets()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "一、项目概述", 1)
    add_para(
        doc,
        "本项目拟建设一套面向人脸三维重建、皮肤纹理分析与可视化交付的软件系统。系统以多视角人脸照片为输入，经过采集管理、图像预处理、三维几何重建、纹理融合、皮肤区域分析、模型展示与结果归档等环节，形成可被甲方查看、复核和后续扩展使用的三维人脸数字资产。",
    )
    add_para(
        doc,
        "本架构书主要服务于甲方沟通和预期工作确认，重点说明系统建设目标、总体架构、数据链路、核心模块、预期交付内容、验收指标与项目风险控制。文档中的数据、示意图和样例输入均以 zyz 数据集为基础，便于甲方围绕真实采集样例理解项目边界。",
    )
    add_table(
        doc,
        ["项目要素", "说明"],
        [
            ["建设目标", "形成从三视图人脸采集到三维模型输出、皮肤纹理分析、Web端预览的完整软件链路。"],
            ["核心输入", "zyz 数据集中的 camera1、camera2、camera3 三张多视角人脸照片。"],
            ["核心输出", "带纹理的三维人脸模型、UV纹理图、阶段过程图、质量指标、演示界面和项目文档。"],
            ["使用对象", "甲方业务人员、项目验收人员、后续系统运维和研发人员。"],
            ["建设原则", "先打通端到端流程，再围绕几何真实度、纹理一致性、展示体验和可验收性持续优化。"],
        ],
        widths=[3.4, 12.2],
    )

    add_heading(doc, "二、甲方关注的预期工作内容", 1)
    add_bullets(
        doc,
        [
            "完成 zyz 三视图样例数据的导入、视角标注、基础质量检查和采集规范沉淀。",
            "建立图像预处理模块，包括人脸区域检测、关键点识别、视角标准化、掩膜生成和异常输入提示。",
            "构建三维几何重建链路，结合 MICA、DECA、FLAME 等模型能力，提升身份形状、侧脸轮廓、鼻部和下颌区域的真实度。",
            "完成深度增强与纹理融合，输出适合展示和归档的三维模型、UV贴图和过程诊断图。",
            "建设后端任务服务，支持上传三视图、创建重建任务、查看进度、下载结果和查询历史记录。",
            "建设 Web 展示界面，支持模型预览、旋转缩放、结果下载和阶段状态展示。",
            "形成可验收的文档体系，包括软件架构书、测试说明、部署说明、数据处理说明和阶段汇报材料。",
        ],
    )

    add_heading(doc, "三、总体架构设计", 1)
    add_para(
        doc,
        "系统采用“数据采集层、算法处理层、服务管理层、前端展示层”的分层架构。采集层负责接收 zyz 这类三视图输入；算法层负责完成预处理、几何重建、纹理融合和皮肤区域分析；服务层通过 FastAPI、SQLite 和 WebSocket 管理任务、状态和结果；展示层通过 Web 页面和 Three.js 模型查看器向甲方呈现最终效果。",
    )
    doc.add_picture(str(ASSETS / "client_architecture.png"), width=Cm(16.2))
    add_caption(doc, "图 2  系统总体架构图")
    add_table(
        doc,
        ["架构层级", "主要职责", "预期成果"],
        [
            ["数据采集层", "接收三视图照片、维护视角关系、记录采集批次和输入质量。", "zyz 数据集样例、采集规范、输入检查记录。"],
            ["算法处理层", "完成关键点、掩膜、几何重建、深度增强、UV纹理融合和皮肤区域处理。", "三维网格、纹理图、诊断图、过程指标。"],
            ["服务管理层", "提供任务创建、进度推送、结果管理、历史查询和接口封装。", "REST API、WebSocket 进度、SQLite任务库。"],
            ["前端展示层", "提供模型预览、状态展示、结果下载和甲方演示入口。", "Web预览页面、GLB加载、结果查看体验。"],
        ],
        widths=[3.0, 7.2, 5.4],
    )

    add_heading(doc, "四、zyz 数据集与数据流", 1)
    add_para(
        doc,
        "zyz 数据集是本架构书中的统一样例数据集，当前包含三张多视角人脸照片，分别对应 camera1、camera2、camera3。系统会将其作为一次完整采集任务处理，并在流程中记录原始输入、预处理结果、重建过程数据和最终输出资产。",
    )
    add_table(
        doc,
        ["数据项", "当前样例", "用途"],
        [
            ["camera1", "camera1_20260511_191357.jpg", "作为多视图重建输入之一，用于补充侧向或偏侧面部信息。"],
            ["camera2", "camera2_20260511_191357.jpg", "作为多视图重建输入之一，通常承担正面或近正面约束。"],
            ["camera3", "camera3_20260511_191357.jpg", "作为多视图重建输入之一，用于补充另一侧面部轮廓与纹理。"],
            ["相机标定", "config/camera_calibration.json", "用于描述相机内参和畸变矫正，提升多视图几何一致性。"],
            ["输出目录", "output/sessions 或 output/meshes", "保存重建模型、纹理、诊断图和任务日志。"],
        ],
        widths=[3.0, 5.0, 7.6],
    )
    doc.add_picture(str(ASSETS / "zyz_data_flow.png"), width=Cm(16.2))
    add_caption(doc, "图 3  zyz 数据驱动的处理流程")

    add_heading(doc, "五、核心功能模块", 1)
    add_table(
        doc,
        ["模块", "建设内容", "甲方可感知成果"],
        [
            ["数据采集与管理", "支持三视图照片上传、视角标注、任务创建、基础质量检查和历史记录。", "可按人员或任务查看输入数据和处理状态。"],
            ["图像预处理", "进行人脸检测、关键点定位、裁剪标准化、掩膜生成、畸变矫正和调试图输出。", "可看到每张输入图是否满足后续重建要求。"],
            ["三维几何重建", "使用 MICA 提供身份形状初始化，使用 DECA 提供姿态与表情估计，并通过 FLAME 多视图优化生成网格。", "输出更贴合被采集对象的三维脸型和侧脸轮廓。"],
            ["深度增强", "引入 Depth-Anything 等深度估计能力，对基础网格进行细节增强和局部形态修正。", "提升面部凹凸、鼻梁、下颌等区域的视觉真实感。"],
            ["纹理融合", "融合三视图颜色信息，生成UV纹理图，并对缺失区域进行补全和平滑。", "输出带真实肤色和纹理的三维模型。"],
            ["皮肤分析", "基于面部区域和纹理图开展肤色、区域纹理、可视化标注等能力建设。", "为后续皮肤状态展示和报告生成打基础。"],
            ["接口与任务服务", "提供创建任务、查询任务、下载模型、进度推送和异常状态管理。", "甲方可通过系统稳定地发起和查看重建任务。"],
            ["Web展示", "基于 Three.js 加载 GLB 模型，支持旋转、缩放、查看和下载。", "甲方可直观看到重建结果并用于演示。"],
        ],
        widths=[3.0, 8.3, 4.3],
    )

    add_heading(doc, "六、接口与系统运行方式", 1)
    add_para(
        doc,
        "系统后端计划采用 FastAPI 提供接口服务，SQLite 保存任务记录，WebSocket 推送重建进度。该设计便于先快速完成单机或局域网演示，也便于后续迁移到服务器部署、增加用户权限或接入甲方已有系统。",
    )
    add_table(
        doc,
        ["接口/能力", "说明", "面向甲方的价值"],
        [
            ["POST /api/sessions", "上传 left/front/right 三视图照片并创建重建任务。", "形成标准化任务入口。"],
            ["GET /api/sessions", "查询历史任务列表和处理状态。", "便于项目演示、验收和追溯。"],
            ["GET /api/sessions/{id}", "查看单个任务的基础信息。", "便于定位某一次采集与重建结果。"],
            ["GET /api/sessions/{id}/model", "下载生成的 GLB 三维模型。", "便于甲方保存、查看和二次使用。"],
            ["POST /api/calibration", "设置或更新相机内参。", "提升多视图几何精度和稳定性。"],
            ["WebSocket /ws/{id}", "实时推送 loading、landmarks、fitting、texture、done 等阶段状态。", "让长耗时任务具备可见进度。"],
        ],
        widths=[4.1, 6.6, 4.9],
    )

    add_heading(doc, "七、预期交付物", 1)
    add_table(
        doc,
        ["交付类别", "交付内容", "验收关注点"],
        [
            ["软件系统", "后端API、前端展示页面、三维模型查看器、任务管理流程。", "能否完成 zyz 数据从上传到结果展示的端到端流程。"],
            ["算法能力", "预处理、三维几何重建、深度增强、UV纹理融合、皮肤区域处理。", "重建结果是否稳定，关键区域是否具备可解释的过程输出。"],
            ["数据资产", "zyz 三视图输入、过程图、模型文件、纹理文件、相机参数和日志。", "数据是否可追溯，过程是否可复核。"],
            ["文档材料", "软件架构书、部署说明、测试记录、验收指标表、阶段汇报图。", "甲方是否能据此完成验收和后续维护。"],
            ["演示成果", "可在浏览器查看的3D模型效果和任务流程演示。", "是否能直观展示系统价值。"],
        ],
        widths=[3.0, 8.1, 4.5],
    )

    add_heading(doc, "八、工作计划与里程碑", 1)
    add_para(
        doc,
        "项目建议按“数据与需求确认、核心重建链路、纹理与皮肤分析、系统集成展示、验收移交”五个阶段推进。每个阶段均保留可展示成果，避免只在最后阶段集中暴露问题。",
    )
    doc.add_picture(str(ASSETS / "work_plan.png"), width=Cm(16.2))
    add_caption(doc, "图 4  预期工作计划与工作量构成")
    add_table(
        doc,
        ["阶段", "主要工作", "阶段成果"],
        [
            ["阶段一：需求与采集确认", "确认 zyz 数据样例、三视图采集规范、输入质量标准、甲方验收口径。", "采集规范、样例数据说明、需求确认记录。"],
            ["阶段二：核心重建链路", "打通预处理、关键点、MICA/DECA/FLAME重建、基础模型输出。", "基础三维模型、关键过程图、初版指标。"],
            ["阶段三：纹理与皮肤分析", "完成多视图纹理融合、UV贴图、皮肤区域处理和视觉优化。", "带纹理模型、UV图、皮肤分析样例。"],
            ["阶段四：系统集成展示", "完成API、任务库、WebSocket进度、Three.js预览和下载入口。", "可演示系统、端到端流程。"],
            ["阶段五：验收与移交", "进行测试、修复、文档整理、部署说明和验收材料输出。", "验收报告、部署说明、最终交付包。"],
        ],
        widths=[3.7, 8.0, 3.9],
    )

    add_heading(doc, "九、验收指标建议", 1)
    add_table(
        doc,
        ["指标类别", "建议指标", "说明"],
        [
            ["流程完整性", "zyz 三视图可完成端到端重建并生成可下载模型。", "这是最基础的系统验收条件。"],
            ["数据可追溯", "每次任务保留输入图、过程图、模型、纹理和日志。", "便于甲方复核问题和阶段成果。"],
            ["几何质量", "正面五官、侧脸轮廓、鼻部、下颌区域具备可接受真实度。", "以视觉评审和重投影误差辅助判断。"],
            ["纹理质量", "UV纹理分辨率不低于2048，主要面部区域无明显错位和大面积缺失。", "保证模型展示效果。"],
            ["系统体验", "任务状态可查看，完成后可在Web端预览和下载。", "保证甲方演示和实际使用体验。"],
            ["稳定性", "异常输入能返回明确提示，正常输入可重复运行。", "降低现场演示和验收风险。"],
        ],
        widths=[3.0, 7.5, 5.1],
    )

    add_heading(doc, "十、风险与保障措施", 1)
    add_table(
        doc,
        ["风险", "影响", "保障措施"],
        [
            ["采集质量不足", "照片模糊、曝光不稳或视角偏差会影响重建效果。", "建立采集规范和质量检查，必要时提示重新采集。"],
            ["侧脸与下颌重建不稳定", "影响三维形态真实度和甲方观感。", "使用MICA身份初始化、多视图优化和诊断图定位问题。"],
            ["纹理拼接错位", "影响模型真实感，尤其在面颊、鼻翼和边界区域。", "保留视角来源图、置信度图和UV诊断图，逐步优化融合策略。"],
            ["模型依赖较重", "MICA、DECA、Depth模型对环境和显存有要求。", "提供CPU/GPU配置说明，保留可降级路径和部署文档。"],
            ["验收口径不清", "可能导致技术成果与甲方预期不一致。", "在阶段一明确输入规范、可视化效果、指标和文档交付范围。"],
        ],
        widths=[4.0, 5.3, 6.3],
    )

    add_heading(doc, "十一、结论", 1)
    add_para(
        doc,
        "本系统的建设重点不是单一算法演示，而是围绕甲方可使用、可查看、可验收的完整软件链路展开。通过以 zyz 数据集为统一样例，项目可以清晰展示从真实三视图输入到三维人脸模型、纹理结果、皮肤分析能力和Web预览的全过程。",
    )
    add_para(
        doc,
        "在预期工作安排上，建议优先完成端到端闭环，再逐步优化几何真实度、纹理质量、皮肤分析能力和系统展示体验。这样既能尽早形成甲方可见成果，也能为后续扩展更多采集对象、更多分析指标和更完整的业务系统打下基础。",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_doc())
